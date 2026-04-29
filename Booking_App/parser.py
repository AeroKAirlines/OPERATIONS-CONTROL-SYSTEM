import openpyxl
import re
import datetime

def parse_flight_data(file_path):
    wb = openpyxl.load_workbook(file_path, data_only=True)
    parsed_data = {}

    for sheet_name in wb.sheetnames:
        if sheet_name.endswith('_보고'):
            dest = sheet_name.split('_')[2]
        elif sheet_name == '보고용':
            dest = 'CJU' 
        else:
            continue
            
        if dest not in parsed_data:
            parsed_data[dest] = []
            
        ws = wb[sheet_name]
        flight_mapping = {}
        is_parsing = False
        current_year = 2026

        for row_idx in range(1, ws.max_row + 1):
            if ws.row_dimensions[row_idx].hidden:
                continue

            cell_b = ws.cell(row=row_idx, column=2).value
            cell_b_str = str(cell_b).strip() if cell_b is not None else ""

            if cell_b_str.startswith('■'):
                is_parsing = True
                flight_mapping = {}
                
                year_match = re.search(r'(\d{4})년', cell_b_str)
                if year_match:
                    current_year = int(year_match.group(1))
                
                current_dir = "왕편"
                current_hub = "청주"
                
                for r in range(row_idx + 1, min(row_idx + 15, ws.max_row + 1)):
                    if ws.row_dimensions[r].hidden: continue
                    
                    scan_b = ws.cell(row=r, column=2).value
                    scan_b_str = str(scan_b).strip() if scan_b is not None else ""
                    
                    is_date_row = False
                    if isinstance(scan_b, datetime.datetime):
                        is_date_row = True
                    elif re.search(r'\d{2}-\d{2}-\d{2}', scan_b_str):
                        is_date_row = True
                    elif re.search(r'^\d{1,2}/\d{1,2}\(', scan_b_str):
                        is_date_row = True
                        
                    if is_date_row:
                        break 
                        
                    for c in range(2, ws.max_column + 1):
                        val = str(ws.cell(row=r, column=c).value or "").strip()
                        
                        if '/' in val:
                            parts = val.split('/')
                            if parts[0] in ['CJJ', 'ICN']:
                                current_dir = "왕편"
                                current_hub = "인천" if parts[0] == 'ICN' else "청주"
                            elif len(parts) > 1 and parts[1] in['CJJ', 'ICN']:
                                current_dir = "복편"
                                current_hub = "인천" if parts[1] == 'ICN' else "청주"
                        elif val == '청주발':
                            current_dir = "왕편"
                            current_hub = "청주"
                        elif val == '제주발':
                            current_dir = "복편"
                            current_hub = "청주"
                        elif val == '인천발':
                            current_dir = "왕편"
                            current_hub = "인천"
                        
                        if val.isdigit() and 2 <= len(val) <= 4:
                            flight_mapping[val] = {
                                "col": c, "dir": current_dir, "hub": current_hub  
                            }
                continue

            if not is_parsing: continue
            if cell_b_str == 'TTL':
                is_parsing = False
                continue

            is_date = False
            date_str = ""
            
            if isinstance(cell_b, datetime.datetime):
                is_date = True
                date_str = cell_b.strftime("%y-%m-%d") 
            else:
                match1 = re.search(r'\d{2}-\d{2}-\d{2}', cell_b_str)
                match2 = re.search(r'^(\d{1,2})/(\d{1,2})\(', cell_b_str)
                
                if match1:
                    is_date = True
                    date_str = match1.group()
                elif match2:
                    is_date = True
                    month = int(match2.group(1))
                    day = int(match2.group(2))
                    year_str = str(current_year)[-2:]
                    date_str = f"{year_str}-{month:02d}-{day:02d}"

            if is_date and flight_mapping:
                daily_info = {"date": date_str, "flights":[]}

                for f_no, info in flight_mapping.items():
                    base_col = info["col"]
                    cfg = ws.cell(row=row_idx, column=base_col).value
                    
                    if cfg is None or str(cfg).strip() == '' or cfg == 'NOOP':
                        continue

                    ttl = ws.cell(row=row_idx, column=base_col + 3).value
                    lf_raw = ws.cell(row=row_idx, column=base_col + 4).value

                    if isinstance(lf_raw, str) and '#' in lf_raw:
                        lf_val = '-'
                    elif isinstance(lf_raw, (int, float)):
                        lf_val = f"{lf_raw * 100:.1f}%"
                    else:
                        lf_val = str(lf_raw).strip() if lf_raw is not None else "0%"

                    daily_info["flights"].append({
                        "flight_no": f_no, "cfg": cfg if cfg is not None else 0,
                        "ttl": ttl if ttl is not None else 0, "lf": lf_val,
                        "direction": info["dir"], "hub": info["hub"]
                    })

                if daily_info["flights"]:
                    parsed_data[dest].append(daily_info)

    return parsed_data